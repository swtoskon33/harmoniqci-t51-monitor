"""
T5.1 monitor: keeps the HarmoniQCI map of open-source QKD tools (the Excel file) up to date.

  python monitor.py [--config config.yaml] [--no-discovery] [--dry-run]

One run:
  1. promotes the candidates an expert marked "Add" into the map (and the scoring sheet);
  2. for every tool in the map reads it where it lives (GitHub, GitLab, the ETSI Forge, Zenodo or
     PyPI: last commit, latest release, license, Docker files, archived, stars), or checks that
     its page is online;
  3. searches GitHub, GitLab, Zenodo and (with a Libraries.io key) PyPI for new tools on ETSI QKD
     and proposes them in "Candidates";
  4. writes the changes into the Excel file, logs each one in "Changes" and writes a report.

It only fills the columns that can be read automatically. The expert columns (category, what it
does, ETSI APIs, test status, findings, priority, scores) are never changed. Tokens, all optional:
GITHUB_TOKEN (higher rate limits, private repositories), GITLAB_TOKEN (gitlab.com), ZENODO_TOKEN,
LIBRARIES_IO_KEY (needed for the search of packages).
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import yaml

from sources import SourceError, Sources, page_status, url_key

SPDX = re.compile(r"^[A-Za-z0-9.+-]+$")
UNKNOWN = {"", "to be checked", "—", "-", "none"}

# columns of the map: the expert's (never written) and the ones this tool fills
COL_NAME, COL_LINK, COL_LANG, COL_LICENSE = "Tool", "Link", "Language", "License"
COL_COMMIT, COL_DOCKER, COL_STATUS = "Last commit", "Docker", "Test status"
AUTO_COLS = ["Latest release", "Archived", "Stars", "Flags", "Last checked"]
CHANGE_COLS = ["Date", "Tool", "Field", "Old value", "New value"]
CAND_COLS = ["Repo", "Link", "Description", "Language", "License", "Stars", "Last push", "Query",
             "First seen", "Decision", "Comment", "Source"]
ADD, ADDED = "Add", "Added"


def unknown(value) -> bool:
    return str(value if value is not None else "").strip().lower() in UNKNOWN


# ----------------------------------------------------------------------------------------- Excel

def headers(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def ensure_cols(ws, names: list[str]) -> dict[str, int]:
    """Adds missing header cells at the end of row 1, styled like the last header."""
    cols = headers(ws)
    last = ws.cell(1, max(cols.values()))
    for name in names:
        if name not in cols:
            col = max(cols.values()) + 1
            cell = ws.cell(1, col, name)
            cell._style = copy.copy(last._style)
            ws.column_dimensions[cell.column_letter].width = 18
            cols[name] = col
    return cols


def sheet(wb, title: str, cols: list[str]):
    if title in wb.sheetnames:
        ws = wb[title]
        ensure_cols(ws, cols)
        return ws
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = "A2"
    return ws


def tool_rows(ws, cols) -> list[int]:
    """Rows of the map that hold a tool (they have a link)."""
    return [r for r in range(2, ws.max_row + 1)
            if str(ws.cell(r, cols[COL_LINK]).value or "").strip().lower().startswith("http")]


@dataclass
class Run:
    today: str
    changes: list[tuple] = field(default_factory=list)       # (tool, field, old, new)
    flags: list[tuple] = field(default_factory=list)         # (tool, text)
    notes: list[str] = field(default_factory=list)           # sources that could not be searched
    keys: set = field(default_factory=set)                   # every key of the tools in the map
    added: list[str] = field(default_factory=list)           # promoted candidates
    proposed: list[dict] = field(default_factory=list)       # new candidates


def set_value(ws, row: int, col: int, value, tool: str, name: str, run: Run, log: bool = True) -> None:
    """Writes a cell and logs the change. A column of the monitor's own that was empty is filled
    without a log entry (the first run fills them all)."""
    old = ws.cell(row, col).value
    if (str(old).strip() if old is not None else "") == (str(value) if value is not None else ""):
        return
    ws.cell(row, col).value = value
    if log and not (name in AUTO_COLS and old in (None, "")):
        run.changes.append((tool, name, "" if old is None else old, value))


# ------------------------------------------------------------------------------------- the steps

def promote(wb, cfg: dict, run: Run) -> None:
    """Candidates marked "Add" become rows of the map and of the scoring sheet."""
    names = cfg["sheets"]
    if names["candidates"] not in wb.sheetnames:
        return
    cand, tools = wb[names["candidates"]], wb[names["tools"]]
    cc, tc = headers(cand), ensure_cols(tools, AUTO_COLS)
    for r in range(2, cand.max_row + 1):
        if str(cand.cell(r, cc["Decision"]).value or "").strip() != ADD:
            continue
        rows = tool_rows(tools, tc)
        at = (rows[-1] if rows else 1) + 1
        tools.insert_rows(at)
        for col in range(1, tools.max_column + 1):
            if at > 2:
                tools.cell(at, col)._style = copy.copy(tools.cell(at - 1, col)._style)
        name = cand.cell(r, cc["Repo"]).value
        numbers = [tools.cell(x, 1).value for x in rows]
        values = {"#": str(max([int(n) for n in numbers if str(n).isdigit()] or [0]) + 1),
                  COL_NAME: name, COL_LINK: cand.cell(r, cc["Link"]).value,
                  "Category / Layer": "To be checked", "What it does": cand.cell(r, cc["Description"]).value,
                  COL_LANG: cand.cell(r, cc["Language"]).value, COL_LICENSE: cand.cell(r, cc["License"]).value
                  or "To be checked", "ETSI APIs": "To be checked", COL_DOCKER: "To be checked",
                  COL_STATUS: "To evaluate", "Test priority": "To be checked"}
        for key, value in values.items():
            if key in tc:
                tools.cell(at, tc[key], value)
        if tools.auto_filter.ref:
            tools.auto_filter.ref = f"A1:{tools.cell(1, tools.max_column).column_letter}{at}"
        add_score_row(wb, names["scores"], name)
        cand.cell(r, cc["Decision"]).value = ADDED
        run.added.append(name)
        run.changes.append((name, "tool", "", "added from the candidates"))


def add_score_row(wb, title: str, name: str) -> None:
    if title not in wb.sheetnames:
        return
    ws = wb[title]
    # rows of tools are the ones with the total formula; notes follow them
    scored = [r for r in range(2, ws.max_row + 1)
              if any(str(ws.cell(r, c).value or "").startswith("=") for c in range(2, ws.max_column + 1))]
    if not scored:
        return
    last = scored[-1]
    row = last + 1
    ws.insert_rows(row)
    ws.cell(row, 1, name)._style = copy.copy(ws.cell(last, 1)._style)
    for col in range(2, ws.max_column + 1):
        src = ws.cell(last, col)
        dst = ws.cell(row, col)
        dst._style = copy.copy(src._style)
        if isinstance(src.value, str) and src.value.startswith("="):
            dst.value = re.sub(rf"(?<![$\w])([A-Z]+){last}\b", rf"\g<1>{row}", src.value)
    # the 1-5 check of the scores covers the new row too
    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(
            " ".join(re.sub(rf"(\D){last}$", rf"\g<1>{row}", str(r)) for r in dv.sqref.ranges))


def update_tools(wb, cfg: dict, src: Sources, run: Run) -> None:
    ws = wb[cfg["sheets"]["tools"]]
    cols = ensure_cols(ws, AUTO_COLS)
    stale = (dt.date.fromisoformat(run.today) - dt.timedelta(days=int(cfg.get("stale_days", 365)))).isoformat()
    for row in tool_rows(ws, cols):
        name = str(ws.cell(row, cols[COL_NAME]).value or ws.cell(row, cols[COL_LINK]).value)
        link = str(ws.cell(row, cols[COL_LINK]).value).strip()
        run.keys.add(url_key(link))
        flags = []
        source = src.for_url(link)
        if source is None:
            status = page_status(link, src.transport)
            if status != "online":
                flags.append(f"page offline: {status}")
            info = None
        else:
            try:
                info = source.tool(link)
                if info is None:
                    flags.append(f"not found on {source.name} (or private and the token has no access)")
            except SourceError as exc:
                info = None
                flags.append(f"not checked: {exc}")
        if info:
            run.keys |= info["keys"]
            if info["url"] and url_key(info["url"]) != url_key(link):
                set_value(ws, row, cols[COL_LINK], info["url"], name, COL_LINK, run)
                run.keys.add(url_key(info["url"]))
                flags.append(f"moved to {info['url']}")
            if info["last_commit"] is not None:
                set_value(ws, row, cols[COL_COMMIT], info["last_commit"], name, COL_COMMIT, run)
            old_release = str(ws.cell(row, cols["Latest release"]).value or "")
            set_value(ws, row, cols["Latest release"], info["release"], name, "Latest release", run)
            if old_release and info["release"] and old_release != info["release"]:
                flags.append(f"new release {info['release']}")
            if info["archived"] is not None:
                set_value(ws, row, cols["Archived"], "Yes" if info["archived"] else "No", name, "Archived", run)
                if info["archived"]:
                    flags.append("archived: no longer maintained")
            if info["stars"] is not None:
                set_value(ws, row, cols["Stars"], info["stars"], name, "Stars", run, log=False)
            if info["activity"] and info["activity"] < stale:
                what = "commit" if info["last_commit"] else "release"
                flags.append(f"no {what} since {info['activity']}")
            lic = str(ws.cell(row, cols[COL_LICENSE]).value or "").strip()
            if info["license"] and (unknown(lic) or (SPDX.match(lic) and lic != info["license"])):
                set_value(ws, row, cols[COL_LICENSE], info["license"], name, COL_LICENSE, run)
                if not unknown(lic):
                    flags.append(f"license changed: {lic} -> {info['license']}")
            if unknown(ws.cell(row, cols[COL_LANG]).value) and info["language"]:
                set_value(ws, row, cols[COL_LANG], info["language"], name, COL_LANG, run)
            docker = str(ws.cell(row, cols[COL_DOCKER]).value or "").strip()
            if info["docker"] and not docker.startswith("Yes"):
                set_value(ws, row, cols[COL_DOCKER], "Yes", name, COL_DOCKER, run)
            elif info["docker"] is False and docker.startswith("Yes"):
                flags.append("no Dockerfile/compose in the root: to be checked")
        set_value(ws, row, cols["Flags"], "; ".join(flags), name, "Flags", run, log=False)
        set_value(ws, row, cols["Last checked"], run.today, name, "", run, log=False)
        run.flags += [(name, f) for f in flags]


def discover(wb, cfg: dict, src: Sources, run: Run) -> None:
    disc = cfg.get("discovery") or {}
    tools = wb[cfg["sheets"]["tools"]]
    tcols = headers(tools)
    known = run.keys | {url_key(tools.cell(r, tcols[COL_LINK]).value) for r in tool_rows(tools, tcols)}
    ws = sheet(wb, cfg["sheets"]["candidates"], CAND_COLS)
    cc = headers(ws)
    seen = {url_key(ws.cell(r, cc["Link"]).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, cc["Link"]).value}
    topic = re.compile(disc.get("topic", r"qkd|quantum[ _-]?key|etsi"), re.I)
    relevant = re.compile(disc.get("relevant", "."), re.I)
    since = (dt.date.fromisoformat(run.today) - dt.timedelta(days=int(disc.get("max_age_days", 730)))).isoformat()
    for source, queries in src.searches(cfg):
        for query in queries:
            try:
                found = source.search(query, int(disc.get("per_query", 30)))
            except SourceError as exc:
                run.notes.append(f"{source.name}, search {query!r}: {exc}")
                break                                          # the source is down: skip its other queries
            for item in found:
                text = " ".join([item["name"], item["description"], *map(str, item["topics"])])
                if (item["key"] in known or item["keys"] & known or item["fork"] or item["archived"]
                        or item["pushed"] < since or not (topic.search(text) and relevant.search(text))):
                    continue
                values = {"Stars": item["stars"], "Last push": item["pushed"]}
                hit = next((seen[k] for k in [item["key"], *item["keys"]] if k in seen), None)
                if hit:                                        # already proposed: refresh only
                    for key, value in values.items():
                        ws.cell(hit, cc[key]).value = value
                    continue
                row = ws.max_row + 1
                entry = {"Repo": item["name"], "Link": item["url"], "Description": item["description"][:300],
                         "Language": item["language"], "License": item["license"], **values, "Query": query,
                         "First seen": run.today, "Decision": "", "Comment": "", "Source": item["source"]}
                for key, value in entry.items():
                    ws.cell(row, cc[key], value)
                seen[item["key"]] = row
                for k in item["keys"]:
                    seen[k] = row
                run.proposed.append(entry)


def log_changes(wb, cfg: dict, run: Run) -> None:
    ws = sheet(wb, cfg["sheets"]["changes"], CHANGE_COLS)
    for tool, name, old, new in run.changes:
        ws.append([run.today, tool, name, old, new])


def report(run: Run) -> str:
    lines = [f"# T5.1 monitor, {run.today}", ""]
    lines += [f"Changes: {len(run.changes)}. Flags: {len(run.flags)}. New candidates: {len(run.proposed)}.", ""]
    if run.added:
        lines += ["## Added from the candidates", ""] + [f"- {n}" for n in run.added] + [""]
    if run.flags:
        lines += ["## Needs expert review", ""] + [f"- **{t}**: {f}" for t, f in run.flags] + [""]
    if run.changes:
        lines += ["## Changes written to the Excel file", "", "| Tool | Field | Old | New |", "| --- | --- | --- | --- |"]
        lines += [f"| {t} | {f} | {o} | {n} |" for t, f, o, n in run.changes] + [""]
    if run.proposed:
        lines += ["## New candidates (decide in the sheet \"Candidates\": Add or Reject)", "",
                  "| Name | Source | Stars | Last update | Description |", "| --- | --- | --- | --- | --- |"]
        lines += [f"| [{p['Repo']}]({p['Link']}) | {p['Source']} | {p['Stars']} | {p['Last push']} | "
                  f"{str(p['Description']).replace('|', '/')[:120]} |" for p in run.proposed] + [""]
    if run.notes:
        lines += ["## Sources not reached", ""] + [f"- {n}" for n in run.notes] + [""]
    if not (run.changes or run.flags or run.proposed):
        lines += ["Nothing changed since the last run.", ""]
    return "\n".join(lines)


def main(argv=None, transport=None, today: str | None = None, sleep=None) -> Run:
    p = argparse.ArgumentParser(prog="python monitor.py", description=__doc__.split("\n\n")[0])
    p.add_argument("--config", default="config.yaml")
    p.add_argument("--no-discovery", action="store_true", help="only update the tools in the map")
    p.add_argument("--dry-run", action="store_true", help="write the report only, not the Excel file")
    a = p.parse_args(argv)
    base = Path(a.config).resolve().parent
    cfg = yaml.safe_load(Path(a.config).read_text())
    excel = base / cfg["excel"]
    wb = openpyxl.load_workbook(excel)
    run = Run(today or dt.date.today().isoformat())
    src = Sources(cfg, os.environ, transport, **({"sleep": sleep} if sleep else {}))
    promote(wb, cfg, run)
    update_tools(wb, cfg, src, run)
    if not a.no_discovery:
        discover(wb, cfg, src, run)
    out = base / cfg.get("report", "reports/latest.md")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(report(run))
    if not a.dry_run:
        log_changes(wb, cfg, run)
        wb.save(excel)
    print(f"{len(run.changes)} changes, {len(run.flags)} flags, {len(run.proposed)} new candidates; report in {out}")
    if os.environ.get("GITHUB_OUTPUT"):                    # for the workflow: open an issue only on news
        with open(os.environ["GITHUB_OUTPUT"], "a") as f:
            f.write(f"news={'true' if run.changes or run.proposed else 'false'}\n")
    return run


if __name__ == "__main__":
    main(sys.argv[1:])
