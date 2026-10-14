"""Tests of monitor.py against fake GitHub, GitLab, Zenodo, PyPI and Libraries.io APIs (no network)."""
import shutil
from pathlib import Path

import httpx
import openpyxl
import pytest
import yaml

import monitor

ROOT = Path(__file__).resolve().parents[1]
TODAY = "2026-10-13"
EXPERT = ["Category / Layer", "What it does", "ETSI APIs", "Test status", "Findings / Notes",
          "Test priority"]


def repo_json(name, **kw):
    return {"html_url": f"https://github.com/{name}", "archived": False, "stargazers_count": 7,
            "language": "Python", "license": {"spdx_id": "MIT"}, "default_branch": "main",
            "pushed_at": "2026-10-01T00:00:00Z", **kw}


class FakeGitHub:
    """Every repository answers with a commit on 2026-10-01 and no release, unless overridden."""

    def __init__(self, repos=None, commits=None, releases=None, files=None, search=None):
        self.repos, self.commits = repos or {}, commits or {}
        self.releases, self.files, self.results = releases or {}, files or {}, search or []
        self.calls = []

    def handler(self, request: httpx.Request) -> httpx.Response:
        path = request.url.path
        self.calls.append(path)
        if path == "/search/repositories":
            return httpx.Response(200, json={"items": self.results})
        parts = path.strip("/").split("/")
        name = "/".join(parts[1:3]).lower()
        if name in self.repos and self.repos[name] is None:
            return httpx.Response(404, json={})
        if len(parts) == 3:
            return httpx.Response(200, json=self.repos.get(name) or repo_json(name))
        what = parts[3]
        if what == "commits":
            date = self.commits.get(name, "2026-10-01")
            return httpx.Response(200, json=[{"commit": {"committer": {"date": f"{date}T10:00:00Z"}}}])
        if what == "releases":
            rel = self.releases.get(name)
            return httpx.Response(200, json=rel) if rel else httpx.Response(404, json={})
        if what == "contents":
            return httpx.Response(200, json=[{"name": n} for n in self.files.get(name, ["README.md"])])
        return httpx.Response(404, json={})


def pages(request):
    return httpx.Response(200 if "qukaydee" in str(request.url) else 503)


@pytest.fixture
def work(tmp_path):
    # the map as it was on 21/09/2026, so that the tests do not depend on later runs
    shutil.copyfile(ROOT / "tests" / "fixtures" / "tools.xlsx", tmp_path / "tools.xlsx")
    cfg = yaml.safe_load((ROOT / "config.yaml").read_text())
    cfg["excel"], cfg["report"] = "tools.xlsx", "report.md"
    cfg["discovery"]["github"] = ["qkd kme"]
    cfg["discovery"]["gitlab"]["queries"] = []
    cfg["discovery"]["zenodo"] = []
    (tmp_path / "config.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True))
    return tmp_path


def run(work, fake, *args, hosts=None):
    """One run; GitHub answers from `fake`, other hosts from `hosts` ({host: handler}), the rest
    like a web page (QuKayDee online, everything else 503)."""
    hosts = {"api.github.com": fake.handler, **(hosts or {})}

    def route(request):
        return hosts.get(request.url.host, pages)(request)
    return monitor.main(["--config", str(work / "config.yaml"), *args], transport=httpx.MockTransport(route),
                        today=TODAY, sleep=lambda s: None)


def rows(path, title="Tool map"):
    ws = openpyxl.load_workbook(path)[title]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def tool(path, name):
    return next(r for r in rows(path) if r["Tool"] and r["Tool"].startswith(name))


def test_updates_automatic_columns_and_keeps_the_experts(work):
    before = tool(work / "tools.xlsx", "qkd_kme_server")
    fake = FakeGitHub(commits={"thomasarmel/qkd_kme_server": "2026-10-05"},
                      files={"thomasarmel/qkd_kme_server": ["Dockerfile", "README.md"]})
    result = run(work, fake, "--no-discovery")
    after = tool(work / "tools.xlsx", "qkd_kme_server")
    assert after["Last commit"] == "2026-10-05"
    assert after["Archived"] == "No" and after["Stars"] == 7 and after["Last checked"] == TODAY
    assert all(after[k] == before[k] for k in EXPERT)
    assert ("qkd_kme_server", "Last commit", "2026-04-01", "2026-10-05") in result.changes
    log = rows(work / "tools.xlsx", "Changes")
    assert {"Date": TODAY, "Tool": "qkd_kme_server", "Field": "Last commit",
            "Old value": "2026-04-01", "New value": "2026-10-05"} in log
    assert "## Changes written to the Excel file" in (work / "report.md").read_text()


def test_fills_unknown_licence_and_docker_but_not_free_text(work):
    fake = FakeGitHub(repos={"qursa-uc3m/qkd-etsi-api-c-wrapper": repo_json(
        "qursa-uc3m/qkd-etsi-api-c-wrapper", license={"spdx_id": "Apache-2.0"})},
        files={"qursa-uc3m/qkd-etsi-api-c-wrapper": ["docker-compose.yml"]})
    run(work, fake, "--no-discovery")
    wrapper = tool(work / "tools.xlsx", "qkd-etsi-api-c-wrapper")
    assert wrapper["License"] == "Apache-2.0" and wrapper["Docker"] == "Yes"
    assert tool(work / "tools.xlsx", "SeQUeNCe")["License"] == "BSD-style (Argonne)"   # expert's wording kept


def test_flags_archived_stale_moved_missing_and_licence_change(work):
    fake = FakeGitHub(
        repos={"tqsd/qunetsim": repo_json("tqsd/QuNetSim", archived=True),
               "open-qkd-network/qkd-net": repo_json("new-owner/qkd-net"),
               "quBIP/etsi-qkd-004".lower(): None,
               "creeppork/next-door-key-simulator": repo_json("CreepPork/next-door-key-simulator",
                                                              license={"spdx_id": "Apache-2.0"})},
        commits={"tqsd/qunetsim": "2023-06-07"},
        files={"creeppork/next-door-key-simulator": ["README.md"]})
    result = run(work, fake, "--no-discovery")
    text = " | ".join(f"{t}: {f}" for t, f in result.flags)
    assert "QuNetSim: archived" in text and "QuNetSim: no commit since 2023-06-07" in text
    assert "moved to https://github.com/new-owner/qkd-net" in text
    assert "QUBIP" in text and "not found" in text
    assert "license changed: MIT -> Apache-2.0" in text
    assert "no Dockerfile/compose in the root" in text
    assert tool(work / "tools.xlsx", "Open QKD")["Link"] == "https://github.com/new-owner/qkd-net"
    assert tool(work / "tools.xlsx", "NetSquid")["Flags"] == "page offline: HTTP 503"
    assert tool(work / "tools.xlsx", "QuKayDee")["Flags"] in ("", None)


def test_new_release_is_flagged_on_the_next_run(work):
    fake = FakeGitHub(releases={"qkdnetsim/qkdnetsim": {"tag_name": "v2.0", "published_at": "2026-01-01T00:00:00Z"}})
    run(work, fake, "--no-discovery")
    fake.releases["qkdnetsim/qkdnetsim"] = {"tag_name": "v2.1", "published_at": "2026-10-10T00:00:00Z"}
    result = run(work, fake, "--no-discovery")
    assert ("QKDNetSim v2.0", "new release v2.1 (2026-10-10)") in result.flags


def test_discovery_proposes_new_repositories_once(work):
    items = [
        {"full_name": "lab/new-kme", "html_url": "https://github.com/lab/new-kme", "description": "ETSI 014 KME",
         "language": "Go", "license": {"spdx_id": "MIT"}, "stargazers_count": 3, "pushed_at": "2026-09-01T00:00:00Z"},
        {"full_name": "thomasarmel/qkd_kme_server", "html_url": "https://github.com/thomasarmel/qkd_kme_server",
         "pushed_at": "2026-09-01T00:00:00Z"},                                         # already in the map
        {"full_name": "x/fork", "html_url": "https://github.com/x/fork", "fork": True, "pushed_at": "2026-09-01T00:00:00Z"},
        {"full_name": "x/old", "html_url": "https://github.com/x/old", "pushed_at": "2020-01-01T00:00:00Z"},
    ]
    fake = FakeGitHub(search=items)
    first = run(work, fake)
    assert [p["Repo"] for p in first.proposed] == ["lab/new-kme"]
    second = run(work, fake)
    assert second.proposed == []
    cands = [r for r in rows(work / "tools.xlsx", "Candidates") if r["Repo"]]
    assert len(cands) == 1 and cands[0]["First seen"] == TODAY and cands[0]["License"] == "MIT"


def test_candidate_marked_for_addition_joins_the_map_and_the_scores(work):
    items = [{"full_name": "lab/new-kme", "html_url": "https://github.com/lab/new-kme", "description": "ETSI 014 KME",
              "language": "Go", "license": {"spdx_id": "MIT"}, "stargazers_count": 3, "pushed_at": "2026-09-01T00:00:00Z"}]
    fake = FakeGitHub(search=items)
    run(work, fake)
    wb = openpyxl.load_workbook(work / "tools.xlsx")
    ws = wb["Candidates"]
    col = [c.value for c in ws[1]].index("Decision") + 1
    ws.cell(2, col).value = "Add"
    wb.save(work / "tools.xlsx")
    result = run(work, fake)
    assert result.added == ["lab/new-kme"]
    new = tool(work / "tools.xlsx", "lab/new-kme")
    assert new["#"] == "15" and new["Test status"] == "To evaluate" and new["Last commit"] == "2026-10-01"
    names = [r["Tool"] for r in rows(work / "tools.xlsx") if r["Tool"]]
    assert names.index("lab/new-kme") == 14 and names[-1].startswith("Source:")   # before the footnote
    scores = openpyxl.load_workbook(work / "tools.xlsx")["Scoring"]
    last = next(r for r in range(2, scores.max_row + 1) if scores.cell(r, 1).value == "lab/new-kme")
    assert last == 15 and scores.cell(16, 1).value is None          # right after QuKayDee, notes below
    assert "B2:G15" in str(scores.data_validations.dataValidation[0].sqref)
    assert f"B{last}*Criteria!$C$2" in scores.cell(last, 8).value and f"COUNT(B{last}:G{last})" in scores.cell(last, 8).value
    assert [r["Decision"] for r in rows(work / "tools.xlsx", "Candidates") if r["Repo"]] == ["Added"]
    assert result.proposed == []


def test_dry_run_leaves_the_excel_file_alone(work):
    before = (work / "tools.xlsx").read_bytes()
    run(work, FakeGitHub(), "--dry-run")
    assert (work / "tools.xlsx").read_bytes() == before and (work / "report.md").exists()


def test_url_key():
    from sources import url_key
    assert url_key("https://github.com/QKDNetSim/qkdnetsim") == "github.com/qkdnetsim/qkdnetsim"
    assert url_key("https://github.com/a/b.git") == url_key("https://github.com/a/b/tree/main") == "github.com/a/b"
    assert url_key("https://gitlab.com/grp/sub/proj/-/tree/main") == "gitlab.com/grp/sub/proj"
    assert url_key("https://zenodo.org/records/123") == url_key("https://doi.org/10.5281/zenodo.123") == "zenodo:123"
    assert url_key("https://pypi.org/project/QKD_Sec/") == "pypi:qkd-sec"
    assert url_key("https://qukaydee.com") == "qukaydee.com"


def add_tools(work, *links):
    """Puts tools into the map the way an expert does: as candidates marked Add."""
    wb = openpyxl.load_workbook(work / "tools.xlsx")
    ws = monitor.sheet(wb, "Candidates", monitor.CAND_COLS)
    cc = monitor.headers(ws)
    for link in links:
        row = ws.max_row + 1
        for key, value in {"Repo": link.rstrip("/").split("/")[-1], "Link": link, "Decision": "Add"}.items():
            ws.cell(row, cc[key], value)
    wb.save(work / "tools.xlsx")


def json_handler(routes):
    """A fake API: {path: JSON} (a path may end with '*' to match a prefix), 404 otherwise."""
    def handler(request):
        raw = request.url.raw_path.decode().split("?")[0]               # keeps %2F of GitLab project paths
        for path, body in routes.items():
            if raw == path or (path.endswith("*") and raw.startswith(path[:-1])):
                return httpx.Response(200, json=body(request) if callable(body) else body)
        return httpx.Response(404, json={})
    return handler


def test_gitlab_and_etsi_forge_tools_are_read(work):
    add_tools(work, "https://gitlab.com/lab/qkd-kms", "https://forge.etsi.org/rep/qkd/gs-qkd-015")
    gitlab = json_handler({
        "/api/v4/projects/lab%2Fqkd-kms": {"web_url": "https://gitlab.com/lab/qkd-kms", "archived": False,
                                           "star_count": 4, "license": {"key": "apache-2.0"},
                                           "last_activity_at": "2026-10-02T00:00:00Z"},
        "/api/v4/projects/lab%2Fqkd-kms/repository/commits": [{"committed_date": "2026-10-02T08:00:00Z"}],
        "/api/v4/projects/lab%2Fqkd-kms/releases": [{"tag_name": "v0.3", "released_at": "2026-09-30T00:00:00Z"}],
        "/api/v4/projects/lab%2Fqkd-kms/repository/tree": [{"name": "Dockerfile"}],
        "/api/v4/projects/lab%2Fqkd-kms/languages": {"Go": 80.0, "Shell": 20.0}})
    forge = json_handler({
        "/rep/api/v4/projects/qkd%2Fgs-qkd-015": {"web_url": "https://forge.etsi.org/rep/qkd/gs-qkd-015",
                                                  "archived": True, "star_count": 1, "license": None},
        "/rep/api/v4/projects/qkd%2Fgs-qkd-015/repository/commits": [{"committed_date": "2024-01-01T00:00:00Z"}],
        "/rep/api/v4/projects/qkd%2Fgs-qkd-015/repository/tree": [{"name": "yang"}]})
    result = run(work, FakeGitHub(), "--no-discovery", hosts={"gitlab.com": gitlab, "forge.etsi.org": forge})
    kms = tool(work / "tools.xlsx", "qkd-kms")
    assert (kms["Last commit"], kms["Latest release"], kms["License"], kms["Docker"], kms["Language"], kms["Stars"]) == \
        ("2026-10-02", "v0.3 (2026-09-30)", "Apache-2.0", "Yes", "Go", 4)
    forge_row = tool(work / "tools.xlsx", "gs-qkd-015")
    assert forge_row["Archived"] == "Yes" and "no commit since 2024-01-01" in forge_row["Flags"]
    assert ("gs-qkd-015", "archived: no longer maintained") in result.flags


def test_zenodo_and_pypi_tools_are_read(work):
    add_tools(work, "https://zenodo.org/records/100", "https://pypi.org/project/qkdsec/")
    zenodo = json_handler({
        "/api/records/100": {"id": 100, "conceptrecid": 99, "links": {"latest": "https://zenodo.org/api/records/105"},
                             "metadata": {"version": "1.0", "publication_date": "2025-01-01"}},
        "/api/records/105": {"id": 105, "conceptrecid": 99, "links": {"latest": "https://zenodo.org/api/records/105"},
                             "metadata": {"version": "1.2", "publication_date": "2026-09-01", "license": {"id": "mit"}}}})
    pypi = json_handler({"/pypi/qkdsec/json": {"info": {"version": "0.4.0", "license_expression": "MIT",
                                                        "project_urls": {"Source": "https://github.com/John-Jepsen/qkdsec"}},
                                               "releases": {"0.4.0": [{"upload_time": "2026-08-20T10:00:00"}]}}})
    run(work, FakeGitHub(), "--no-discovery", hosts={"zenodo.org": zenodo, "pypi.org": pypi})
    z = tool(work / "tools.xlsx", "100")
    assert (z["Latest release"], z["License"], z["Last commit"], z["Archived"]) == ("1.2 (2026-09-01)", "MIT", None, None)
    p = tool(work / "tools.xlsx", "qkdsec")
    assert (p["Latest release"], p["License"], p["Language"]) == ("0.4.0 (2026-08-20)", "MIT", "Python")


def test_discovery_on_gitlab_zenodo_and_librariesio_skips_what_the_map_has(work, monkeypatch):
    cfg = yaml.safe_load((work / "config.yaml").read_text())
    cfg["discovery"]["gitlab"] = {"hosts": ["https://gitlab.com"], "queries": ["qkd"]}
    cfg["discovery"]["zenodo"] = ["QKD"]
    (work / "config.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True))
    monkeypatch.setenv("LIBRARIES_IO_KEY", "k")
    gitlab = json_handler({"/api/v4/projects": [
        {"path_with_namespace": "lab/qkd-sdn-agent", "web_url": "https://gitlab.com/lab/qkd-sdn-agent",
         "description": "SDN agent for QKD networks", "star_count": 2, "last_activity_at": "2026-09-01T00:00:00Z"},
        {"path_with_namespace": "x/qkd-fork", "web_url": "https://gitlab.com/x/qkd-fork", "description": "QKD KMS",
         "forked_from_project": {"id": 1}, "last_activity_at": "2026-09-01T00:00:00Z"}]})
    zenodo = json_handler({"/api/records": {"hits": {"hits": [
        {"id": 7, "conceptrecid": 6, "metadata": {"title": "ETSI QKD 014 KME emulator", "publication_date": "2026-05-01",
                                                  "description": "<p>A KME for ETSI GS QKD 014</p>", "license": {"id": "mit"}}},
        {"id": 9, "conceptrecid": 8, "metadata": {
            "title": "qkd_kme_server v1.0", "publication_date": "2026-05-01", "description": "QKD KME server",
            "related_identifiers": [{"identifier": "https://github.com/thomasarmel/qkd_kme_server/tree/v1.0",
                                     "relation": "isSupplementTo"}]}}]}}})
    libraries = json_handler({"/api/search": [
        {"name": "etsi-qkd-kms", "platform": "PyPI", "description": "ETSI QKD KMS client",
         "package_manager_url": "https://pypi.org/project/etsi-qkd-kms/", "stars": 1,
         "latest_release_published_at": "2026-07-01T00:00:00Z"}]})
    result = run(work, FakeGitHub(), hosts={"gitlab.com": gitlab, "zenodo.org": zenodo, "libraries.io": libraries})
    assert sorted((p["Repo"], p["Source"]) for p in result.proposed) == [
        ("ETSI QKD 014 KME emulator", "Zenodo"), ("etsi-qkd-kms (PyPI)", "Libraries.io (PyPI)"),
        ("lab/qkd-sdn-agent", "GitLab (gitlab.com)")]
    zen = next(p for p in result.proposed if p["Source"] == "Zenodo")
    assert zen["Link"] == "https://zenodo.org/records/6" and zen["Description"] == "A KME for ETSI GS QKD 014"
    assert run(work, FakeGitHub(), hosts={"gitlab.com": gitlab, "zenodo.org": zenodo, "libraries.io": libraries}).proposed == []


def test_a_source_that_is_down_is_reported_and_the_run_goes_on(work):
    cfg = yaml.safe_load((work / "config.yaml").read_text())
    cfg["discovery"]["zenodo"] = ["QKD", "ETSI QKD"]
    (work / "config.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True))
    add_tools(work, "https://zenodo.org/records/100")

    def down(request):
        raise httpx.ConnectError("no route", request=request)
    result = run(work, FakeGitHub(), hosts={"zenodo.org": down})
    assert result.notes == ["Zenodo, search 'QKD': Zenodo not reachable (ConnectError)"]
    assert ("100", "not checked: Zenodo not reachable (ConnectError)") in result.flags
    assert "## Sources not reached" in (work / "report.md").read_text()
    assert tool(work / "tools.xlsx", "qkd_kme_server")["Last commit"] == "2026-10-01"     # GitHub still read
